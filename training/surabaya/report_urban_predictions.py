"""Generate an Excel prediction-vs-groundtruth report for Surabaya baselines.

The workbook mirrors the prediction-vs-groundtruth part of the Dhompo xls_15
report at a smaller scope: metrics summary and best-model per-horizon sheets.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "src"))

import joblib
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.preprocessing import StandardScaler

from dhompo.config import load_yaml_config, resolve_artifact_path
from dhompo.data.urban_features import (
    align_urban_features_targets,
    build_urban_delta_targets,
    build_urban_forecast_features,
    build_urban_targets,
)
from dhompo.data.urban_loader import preprocess_urban_wide_data
from training.evaluate import calc_metrics, performance_grade


DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "reports"
    / "surabaya"
    / "tables"
    / "xls_17_surabaya_residual_vs_persistence_ringkas.xlsx"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create Surabaya prediction report workbook")
    parser.add_argument(
        "--config",
        default="configs/surabaya/urban_water_level.yaml",
        help="Surabaya/urban config path",
    )
    parser.add_argument(
        "--model-dir",
        default="models/surabaya/urban",
        help="Directory containing training_metadata.json and model pkl files",
    )
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT),
        help="Output xlsx path",
    )
    parser.add_argument(
        "--no-persistence-baselines",
        action="store_true",
        help="Exclude Persistence and Trend Persistence rows from Ringkasan and horizon sheets.",
    )
    return parser.parse_args()


def _load_metadata(model_dir: Path) -> dict[str, Any]:
    metadata_path = model_dir / "training_metadata.json"
    if not metadata_path.exists():
        raise FileNotFoundError(f"Urban metadata not found: {metadata_path}")
    return json.loads(metadata_path.read_text(encoding="utf-8"))


def _model_display_name(model_key: str) -> str:
    return {
        "ridge": "Ridge",
        "gradient_boosting": "Gradient Boosting",
        "xgboost": "XGBoost",
        "random_forest": "Random Forest",
        "lasso": "Lasso",
        "linear_regression": "Linear Regression",
        "elasticnet": "ElasticNet",
    }.get(model_key, model_key.replace("_", " ").title())


def _read_model_path(model_dir: Path, raw_path: str) -> Path:
    return resolve_artifact_path(model_dir, raw_path)


def _feature_list_frame(feature_columns: list[str]) -> pd.DataFrame:
    rows = []
    for i, feature in enumerate(feature_columns, start=1):
        transform = "t0"
        base = feature
        if "_flag_missing_lag" in feature:
            base, lag = feature.rsplit("_flag_missing_lag", 1)
            transform = f"flag_missing_lag{lag}"
        elif "_flag_outlier_lag" in feature:
            base, lag = feature.rsplit("_flag_outlier_lag", 1)
            transform = f"flag_outlier_lag{lag}"
        elif "_flag_stale_lag" in feature:
            base, lag = feature.rsplit("_flag_stale_lag", 1)
            transform = f"flag_stale_lag{lag}"
        elif feature.endswith("_flag_missing"):
            base = feature.removesuffix("_flag_missing")
            transform = "flag_missing"
        elif feature.endswith("_flag_outlier"):
            base = feature.removesuffix("_flag_outlier")
            transform = "flag_outlier"
        elif feature.endswith("_flag_stale"):
            base = feature.removesuffix("_flag_stale")
            transform = "flag_stale"
        elif "_rmean_" in feature:
            base, label = feature.rsplit("_rmean_", 1)
            transform = f"rolling_mean_{label}"
        elif "_rstd_" in feature:
            base, label = feature.rsplit("_rstd_", 1)
            transform = f"rolling_std_{label}"
        elif feature.endswith("_diff1"):
            base = feature.removesuffix("_diff1")
            transform = "diff1"
        elif feature.endswith("_diff2"):
            base = feature.removesuffix("_diff2")
            transform = "diff2"
        elif "_lag" in feature:
            base, lag = feature.rsplit("_lag", 1)
            transform = f"lag{lag}"
        elif feature.endswith("_t0"):
            base = feature.removesuffix("_t0")
            transform = "t0"
        elif feature in {"hour_sin", "hour_cos", "dayofweek", "is_night"}:
            base = "datetime"
            transform = feature

        rows.append(
            {
                "No": i,
                "Feature": feature,
                "Sinyal": base,
                "Transformasi": transform,
                "Keterangan": "Urban baseline feature",
            }
        )
    return pd.DataFrame(rows)


def _descriptive_stats_frame(data) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for column in data.canonical.columns:
        series = data.canonical[column].dropna()
        diffs = series.diff().abs().dropna()
        flags = data.quality_flags[column].value_counts()
        coverage_row = data.coverage[data.coverage["column"] == column]
        coverage = (
            float(coverage_row["coverage"].iloc[0])
            if not coverage_row.empty
            else float(series.notna().mean())
        )

        row = {
            "Kolom": column,
            "N_Observed": int(series.count()),
            "N_Total": int(len(data.canonical)),
            "Coverage (%)": round(coverage * 100, 4),
            "Missing_Flag_Count": int(flags.get("MISSING", 0)),
            "Stale_Flag_Count": int(flags.get("STALE", 0)),
            "Outlier_Flag_Count": int(flags.get("OUTLIER", 0)),
            "OK_Flag_Count": int(flags.get("OK", 0)),
            "Mean (cm)": round(float(series.mean()), 4) if not series.empty else None,
            "Std (cm)": round(float(series.std()), 4) if len(series) > 1 else None,
            "Min (cm)": round(float(series.min()), 4) if not series.empty else None,
            "P01 (cm)": round(float(series.quantile(0.01)), 4) if not series.empty else None,
            "P05 (cm)": round(float(series.quantile(0.05)), 4) if not series.empty else None,
            "P25 (cm)": round(float(series.quantile(0.25)), 4) if not series.empty else None,
            "Median (cm)": round(float(series.median()), 4) if not series.empty else None,
            "P75 (cm)": round(float(series.quantile(0.75)), 4) if not series.empty else None,
            "P95 (cm)": round(float(series.quantile(0.95)), 4) if not series.empty else None,
            "P99 (cm)": round(float(series.quantile(0.99)), 4) if not series.empty else None,
            "Max (cm)": round(float(series.max()), 4) if not series.empty else None,
            "Median Abs Diff 30m (cm)": (
                round(float(diffs.median()), 4) if not diffs.empty else None
            ),
            "P95 Abs Diff 30m (cm)": (
                round(float(diffs.quantile(0.95)), 4) if not diffs.empty else None
            ),
            "P99 Abs Diff 30m (cm)": (
                round(float(diffs.quantile(0.99)), 4) if not diffs.empty else None
            ),
            "Max Abs Diff 30m (cm)": (
                round(float(diffs.max()), 4) if not diffs.empty else None
            ),
            "Abs Diff > 50cm Count": int((diffs > 50).sum()),
            "Abs Diff > 100cm Count": int((diffs > 100).sum()),
            "Abs Diff > 200cm Count": int((diffs > 200).sum()),
        }
        rows.append(row)
    return pd.DataFrame(rows)


def _quality_flag_explanation_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Flag": "OK",
                "Arti": "Nilai tersedia sebagai observasi asli pada timestamp tersebut.",
                "Dampak ke fitur": "Dipakai langsung sebagai nilai sensor.",
            },
            {
                "Flag": "STALE",
                "Arti": (
                    "Nilai asli kosong, tetapi masih dapat diisi dari observasi terakhir "
                    "dalam batas max_ffill_steps."
                ),
                "Dampak ke fitur": (
                    "Nilai forward-fill dipakai, dan fitur flag_stale bernilai 1 agar "
                    "model tahu data bukan observasi baru."
                ),
            },
            {
                "Flag": "MISSING",
                "Arti": (
                    "Nilai asli kosong dan tidak ada observasi terakhir yang masih valid "
                    "untuk forward-fill terbatas."
                ),
                "Dampak ke fitur": (
                    "Nilai numerik fitur diisi sentinel 0.0, dan fitur flag_missing "
                    "bernilai 1 agar model tahu sensor tidak tersedia."
                ),
            },
            {
                "Flag": "OUTLIER",
                "Arti": (
                    "Nilai target tersedia di data mentah, tetapi gagal aturan cleaning "
                    "misalnya 0/reset, terlalu tinggi, atau lonjakan 30 menit terlalu besar."
                ),
                "Dampak ke fitur": (
                    "Nilai tersebut tidak dipakai sebagai label training; untuk fitur real-time "
                    "dipakai nilai forward-fill terbatas dan flag_outlier menjaga jejaknya."
                ),
            },
        ]
    )


def _prepare_frames(config_path: str, metadata: dict[str, Any]):
    config = load_yaml_config(config_path)
    data = preprocess_urban_wide_data(config_path=config_path)
    feature_cfg = config.get("features", {})
    rolling_windows = [
        (int(w["steps"]), str(w["label"]))
        for w in feature_cfg.get("rolling_windows", [])
    ]
    X_features = build_urban_forecast_features(
        data.values,
        feature_columns=data.feature_columns,
        quality_flags=data.quality_flags,
        include_quality_flags=bool(feature_cfg.get("include_quality_flags", True)),
        lag_steps=[int(v) for v in feature_cfg.get("lag_steps", [1, 2, 3])],
        rolling_windows=rolling_windows,
    )
    horizons = [int(h) for h in config.get("horizons", [1, 2, 3, 4, 5])]
    future_targets = build_urban_targets(data.modeling_canonical, data.target_column, horizons)
    target_mode = str(metadata.get("target_mode", config.get("target_mode", "level"))).lower()
    if target_mode in {"delta", "persistence_residual"}:
        y_model = build_urban_delta_targets(
            data.modeling_canonical,
            current_values=data.values,
            target_column=data.target_column,
            horizons=horizons,
        )
    else:
        y_model = future_targets
    X_full, y_model = align_urban_features_targets(X_features, y_model)
    future_targets = {h: y.loc[X_full.index] for h, y in future_targets.items()}
    current_target = data.values[data.target_column].loc[X_full.index]
    lag_1h_target = data.values[data.target_column].shift(2).loc[X_full.index]

    train_split = float(config.get("train_split", metadata.get("train_split", 0.8)))
    split_idx = int(len(X_full) * train_split)
    X_test = X_full.iloc[split_idx:][metadata["feature_columns"]]
    y_test = {h: y.iloc[split_idx:] for h, y in future_targets.items()}
    current_test = current_target.iloc[split_idx:]
    trend_1h_test = (current_target - lag_1h_target).iloc[split_idx:].fillna(0.0)
    return data, X_full, X_test, y_test, horizons, current_test, trend_1h_test


def _append_output_and_summary(
    all_rows: list[pd.DataFrame],
    summary_rows: list[dict[str, Any]],
    *,
    scenario: str,
    horizon: int,
    model_name: str,
    index: pd.DatetimeIndex,
    obs: pd.Series,
    pred: pd.Series,
) -> None:
    output = pd.DataFrame(
        {
            "Datetime": index,
            "Target Datetime": index + pd.Timedelta(hours=horizon),
            "Horizon": f"+{horizon} Jam",
            "Skenario": scenario,
            "Model": model_name,
            "Observasi (cm)": obs.to_numpy(),
            "Prediksi (cm)": pred.to_numpy(),
        }
    )
    output["Error (cm)"] = output["Observasi (cm)"] - output["Prediksi (cm)"]
    output["Abs Error (cm)"] = output["Error (cm)"].abs()
    numeric_cols = [
        "Observasi (cm)",
        "Prediksi (cm)",
        "Error (cm)",
        "Abs Error (cm)",
    ]
    output[numeric_cols] = output[numeric_cols].round(4)
    all_rows.append(output)

    metrics = calc_metrics(obs.to_numpy(), pred.to_numpy())
    summary_rows.append(
        {
            "Skenario": scenario,
            "Horizon": f"+{horizon} Jam",
            "Model": model_name,
            "NSE": round(metrics["NSE"], 4),
            "RMSE (cm)": round(metrics["RMSE"], 4),
            "MAE (cm)": round(metrics["MAE"], 4),
            "R2": round(metrics["R2"], 4),
            "PBIAS (%)": round(metrics["PBIAS"], 4),
            "Grade": performance_grade(metrics["NSE"]),
            "N_Samples": len(output),
        }
    )


def _prediction_outputs(
    model_dir: Path,
    metadata: dict[str, Any],
    X_test: pd.DataFrame,
    y_test: dict[int, pd.Series],
    current_test: pd.Series,
    trend_1h_test: pd.Series,
    include_persistence_baselines: bool = True,
) -> tuple[pd.DataFrame, dict[int, pd.DataFrame], pd.DataFrame]:
    scaler: StandardScaler = joblib.load(model_dir / "standard_scaler_global.pkl")
    X_scaled = pd.DataFrame(
        scaler.transform(X_test),
        columns=X_test.columns,
        index=X_test.index,
    )
    scenario = "Surabaya: Baseline 30-menit"

    all_rows = []
    summary_rows = []
    best_sheets: dict[int, pd.DataFrame] = {}
    target_mode = str(metadata.get("target_mode", "level")).lower()

    if include_persistence_baselines:
        for horizon, obs in y_test.items():
            persistence = pd.Series(current_test.to_numpy(), index=X_test.index)
            trend_persistence = pd.Series(
                current_test.to_numpy() + trend_1h_test.to_numpy() * horizon,
                index=X_test.index,
            )
            _append_output_and_summary(
                all_rows,
                summary_rows,
                scenario=scenario,
                horizon=horizon,
                model_name="Persistence",
                index=X_test.index,
                obs=obs,
                pred=persistence,
            )
            _append_output_and_summary(
                all_rows,
                summary_rows,
                scenario=scenario,
                horizon=horizon,
                model_name="Trend Persistence",
                index=X_test.index,
                obs=obs,
                pred=trend_persistence,
            )

    for model_id, model_meta in metadata["models"].items():
        horizon_text, model_key = model_id.split(":", 1)
        horizon = int(horizon_text.removeprefix("h"))
        model_path = _read_model_path(model_dir, model_meta["path"])
        model = joblib.load(model_path)
        X = X_scaled if model_meta["use_scaled"] else X_test
        raw_pred = model.predict(X)
        if target_mode in {"delta", "persistence_residual"}:
            raw_pred = current_test.to_numpy() + raw_pred
        pred = pd.Series(raw_pred, index=X_test.index, name="Prediksi (cm)")
        obs = y_test[horizon].reindex(X_test.index)
        _append_output_and_summary(
            all_rows,
            summary_rows,
            scenario=scenario,
            horizon=horizon,
            model_name=_model_display_name(model_key),
            index=X_test.index,
            obs=obs,
            pred=pred,
        )

    all_output = pd.concat(all_rows, ignore_index=True)
    summary = pd.DataFrame(summary_rows).sort_values(["Horizon", "NSE"], ascending=[True, False])

    for horizon in y_test:
        best_row = summary[summary["Horizon"] == f"+{horizon} Jam"].iloc[0]
        model_name = str(best_row["Model"])
        best_sheets[horizon] = all_output[
            (all_output["Horizon"] == f"+{horizon} Jam")
            & (all_output["Model"] == model_name)
        ].reset_index(drop=True)

    return all_output, best_sheets, summary


def _autosize_and_style(path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 35)
    wb.save(path)


def main() -> None:
    args = parse_args()
    model_dir = Path(args.model_dir)
    if not model_dir.is_absolute():
        model_dir = PROJECT_ROOT / model_dir
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = PROJECT_ROOT / output_path
    output_path.parent.mkdir(parents=True, exist_ok=True)

    metadata = _load_metadata(model_dir)
    data, _, X_test, y_test, horizons, current_test, trend_1h_test = _prepare_frames(
        args.config,
        metadata,
    )
    _, best_sheets, summary = _prediction_outputs(
        model_dir,
        metadata,
        X_test,
        y_test,
        current_test,
        trend_1h_test,
        include_persistence_baselines=not args.no_persistence_baselines,
    )
    descriptive_stats = _descriptive_stats_frame(data)
    flag_explanations = _quality_flag_explanation_frame()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Ringkasan", index=False)
        for horizon in horizons:
            best_sheets[horizon].to_excel(
                writer,
                sheet_name=f"Surabaya_Baseline_h{horizon}",
                index=False,
            )
        descriptive_stats.to_excel(
            writer,
            sheet_name="Statistika_Deskriptif",
            index=False,
        )
        flag_explanations.to_excel(
            writer,
            sheet_name="Penjelasan_Flag",
            index=False,
        )

    _autosize_and_style(output_path)
    print(f"Report written: {output_path}")
    print(f"Rows: test_samples={len(X_test)}")


if __name__ == "__main__":
    main()
