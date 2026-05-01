"""Export prediksi vs ground truth untuk Skenario A, B, C ke satu file Excel.

Skenario:
  A = Baseline Train 2022 -> Test 2023
  B = Combined Training (2022 + 80% 2023) -> Test 20% akhir 2023
  C = Combined Training + Rainfall features -> Test 20% akhir 2023

Output:
  reports/tables/xls_14_prediksi_vs_groundtruth_skenarioABC.xlsx

Struktur sheet:
  - Data_Mentah_2022      : raw 2022
  - Data_Mentah_2023      : raw 2023 (stations + rainfall)
  - A_Baseline2022_hX     : tabel datetime, obs, pred, error (+ chart)
  - B_Combined_hX         : idem
  - C_CombinedRain_hX     : idem
  - Ringkasan             : NSE/RMSE/MAE/R2/PBIAS per skenario per horizon
"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from sklearn.base import clone
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor
from sklearn.linear_model import Lasso, LinearRegression, Ridge
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.preprocessing import StandardScaler
from xgboost import XGBRegressor

from dhompo.data.features import (
    align_features_targets,
    build_features_from_segments,
    build_forecast_features,
    build_targets,
)
from dhompo.data.loader import (
    ALL_STATIONS,
    RAINFALL_COLUMN,
    TARGET_STATION,
    UPSTREAM_STATIONS,
    load_combined_data,
    load_data,
    load_generated_data,
)

DATA_CLEAN = PROJECT_ROOT / "data" / "data-clean.csv"
DATA_GEN = PROJECT_ROOT / "data" / "Data generated 2023.xlsx"
OUT_XLSX = PROJECT_ROOT / "reports" / "tables" / "xls_14_prediksi_vs_groundtruth_skenarioABC.xlsx"

HORIZONS = [1, 2, 3, 4, 5]
HORIZON_STEPS = {h: h * 2 for h in HORIZONS}

MODEL_LIBRARY = {
    "Linear Regression": (LinearRegression(), True),
    "Ridge Regression": (Ridge(alpha=1.0), True),
    "Lasso (L1)": (Lasso(alpha=0.01, max_iter=10000), True),
    "Random Forest": (
        RandomForestRegressor(n_estimators=200, max_depth=12, random_state=42, n_jobs=-1),
        False,
    ),
    "Gradient Boosting": (
        GradientBoostingRegressor(n_estimators=200, max_depth=5, learning_rate=0.1, random_state=42),
        False,
    ),
    "XGBoost": (
        XGBRegressor(
            n_estimators=200, max_depth=5, learning_rate=0.1, random_state=42, n_jobs=-1, verbosity=0
        ),
        False,
    ),
}


def hydro_metrics(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
    rmse = float(np.sqrt(mean_squared_error(y_true, y_pred)))
    mae = float(mean_absolute_error(y_true, y_pred))
    r2 = float(r2_score(y_true, y_pred))
    nse = float(1 - np.sum((y_true - y_pred) ** 2) / np.sum((y_true - np.mean(y_true)) ** 2))
    pbias = float(100 * np.sum(y_pred - y_true) / np.sum(y_true))
    return {"NSE": nse, "RMSE": rmse, "MAE": mae, "R2": r2, "PBIAS": pbias}


def run_scenario_A(df_clean, df_gen, common_stations):
    print("[A] Train 2022 -> Test 2023 ...")
    X_2022 = build_forecast_features(df_clean, UPSTREAM_STATIONS, target=TARGET_STATION)
    y_2022_h = build_targets(df_clean, HORIZONS, HORIZON_STEPS, target=TARGET_STATION)
    X_2022, y_2022_h = align_features_targets(X_2022, y_2022_h)

    gen_common = df_gen[common_stations].copy()
    X_2023 = build_forecast_features(gen_common, UPSTREAM_STATIONS, target=TARGET_STATION)
    y_2023_h = build_targets(gen_common, HORIZONS, HORIZON_STEPS, target=TARGET_STATION)
    X_2023, y_2023_h = align_features_targets(X_2023, y_2023_h)

    scaler = StandardScaler()
    Xtr_std = pd.DataFrame(scaler.fit_transform(X_2022), index=X_2022.index, columns=X_2022.columns)
    Xte_std = pd.DataFrame(scaler.transform(X_2023), index=X_2023.index, columns=X_2023.columns)

    horizon_best = {}
    for h in HORIZONS:
        best = None
        for name, (tmpl, use_std) in MODEL_LIBRARY.items():
            model = clone(tmpl)
            Xtr = Xtr_std if use_std else X_2022
            Xte = Xte_std if use_std else X_2023
            model.fit(Xtr, y_2022_h[h])
            y_hat = model.predict(Xte)
            m = hydro_metrics(y_2023_h[h].values, y_hat)
            if best is None or m["NSE"] > best["metrics"]["NSE"]:
                best = {"model": name, "metrics": m, "y_true": y_2023_h[h], "y_pred": y_hat}
        print(f"  h{h}: {best['model']}  NSE={best['metrics']['NSE']:.4f}")
        horizon_best[h] = best
    return horizon_best


def run_scenario_combined(with_rainfall: bool):
    label = "C" if with_rainfall else "B"
    print(f"[{label}] Combined training{' + rainfall' if with_rainfall else ''} ...")
    segments = load_combined_data(DATA_CLEAN, DATA_GEN)

    extra = [RAINFALL_COLUMN] if with_rainfall else None
    X_all = build_features_from_segments(
        segments, upstream_stations=UPSTREAM_STATIONS, target=TARGET_STATION, extra_columns=extra
    )

    y_parts = []
    for seg in segments:
        y_seg = build_targets(seg.df, HORIZONS, HORIZON_STEPS, target=TARGET_STATION)
        y_parts.append(pd.concat({h: s for h, s in y_seg.items()}, axis=1))
    y_all = pd.concat(y_parts, axis=0)
    y_horizons = {h: y_all[h] for h in HORIZONS}

    X_all, y_horizons = align_features_targets(X_all, y_horizons)

    seg_2023_start = segments[1].df.index.min()
    idx_2023 = X_all.index >= seg_2023_start
    n_2022 = int((~idx_2023).sum())
    n_2023 = int(idx_2023.sum())
    split = n_2022 + int(n_2023 * 0.8)

    X_train = X_all.iloc[:split]
    X_test = X_all.iloc[split:]
    scaler = StandardScaler()
    Xtr_std = pd.DataFrame(scaler.fit_transform(X_train), index=X_train.index, columns=X_train.columns)
    Xte_std = pd.DataFrame(scaler.transform(X_test), index=X_test.index, columns=X_test.columns)

    horizon_best = {}
    for h in HORIZONS:
        y_train = y_horizons[h].iloc[:split]
        y_test = y_horizons[h].iloc[split:]
        best = None
        for name, (tmpl, use_std) in MODEL_LIBRARY.items():
            model = clone(tmpl)
            Xtr = Xtr_std if use_std else X_train
            Xte = Xte_std if use_std else X_test
            model.fit(Xtr, y_train)
            y_hat = model.predict(Xte)
            m = hydro_metrics(y_test.values, y_hat)
            if best is None or m["NSE"] > best["metrics"]["NSE"]:
                best = {"model": name, "metrics": m, "y_true": y_test, "y_pred": y_hat}
        print(f"  h{h}: {best['model']}  NSE={best['metrics']['NSE']:.4f}")
        horizon_best[h] = best
    return horizon_best


def pred_frame(best: dict, horizon: int, scenario: str) -> pd.DataFrame:
    y_true = best["y_true"]
    y_pred = np.asarray(best["y_pred"])
    df = pd.DataFrame(
        {
            "Datetime": y_true.index,
            "Horizon": f"+{horizon} Jam",
            "Skenario": scenario,
            "Model": best["model"],
            "Observasi (m)": np.round(y_true.values, 4),
            "Prediksi (m)": np.round(y_pred, 4),
            "Error (m)": np.round(y_true.values - y_pred, 4),
            "Abs Error (m)": np.round(np.abs(y_true.values - y_pred), 4),
        }
    )
    return df


def add_chart(ws, n_rows: int, title: str):
    from openpyxl.chart import LineChart, Reference

    chart = LineChart()
    chart.title = title
    chart.y_axis.title = "Water level (m)"
    chart.x_axis.title = "Datetime"
    chart.height = 10
    chart.width = 22
    data = Reference(ws, min_col=5, max_col=6, min_row=1, max_row=n_rows + 1)
    cats = Reference(ws, min_col=1, max_col=1, min_row=2, max_row=n_rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "J2")


def main():
    print("Memuat dataset ...")
    df_clean = load_data(DATA_CLEAN)
    gen = load_generated_data(DATA_GEN)
    df_gen = gen.stations
    rainfall = gen.rainfall
    common_stations = [c for c in ALL_STATIONS if c in df_clean.columns and c in df_gen.columns]

    best_A = run_scenario_A(df_clean, df_gen, common_stations)
    best_B = run_scenario_combined(with_rainfall=False)
    best_C = run_scenario_combined(with_rainfall=True)

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    print(f"Menulis Excel: {OUT_XLSX}")

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        raw_2022 = df_clean.copy()
        raw_2022.index.name = "Datetime"
        raw_2022.reset_index().to_excel(writer, sheet_name="Data_Mentah_2022", index=False)

        raw_2023 = df_gen.copy()
        raw_2023[RAINFALL_COLUMN] = rainfall
        raw_2023.index.name = "Datetime"
        raw_2023.reset_index().to_excel(writer, sheet_name="Data_Mentah_2023", index=False)

        summary_rows = []
        scenarios = [
            ("A_Baseline2022", "A: Baseline Train 2022 -> Test 2023", best_A),
            ("B_Combined", "B: Combined Training (2022+2023)", best_B),
            ("C_CombinedRain", "C: Combined + Rainfall", best_C),
        ]
        for prefix, label, best_map in scenarios:
            for h in HORIZONS:
                df = pred_frame(best_map[h], h, label)
                sheet = f"{prefix}_h{h}"[:31]
                df.to_excel(writer, sheet_name=sheet, index=False)
                ws = writer.sheets[sheet]
                add_chart(ws, len(df), f"{label} | Horizon +{h} Jam | {best_map[h]['model']}")
                m = best_map[h]["metrics"]
                summary_rows.append(
                    {
                        "Skenario": label,
                        "Horizon": f"+{h} Jam",
                        "Model": best_map[h]["model"],
                        "NSE": round(m["NSE"], 4),
                        "RMSE (m)": round(m["RMSE"], 4),
                        "MAE (m)": round(m["MAE"], 4),
                        "R2": round(m["R2"], 4),
                        "PBIAS (%)": round(m["PBIAS"], 4),
                        "N_Samples": len(best_map[h]["y_true"]),
                    }
                )
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Ringkasan", index=False)

    print(f"Selesai -> {OUT_XLSX}")


if __name__ == "__main__":
    main()
